import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_sop_document():
    doc = Document()
    
    # Page Setup (A4 Portrait, 2 cm margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    # Colors
    c_primary = RGBColor(13, 110, 85)     # Palm Green (#0D6E55)
    c_dark = RGBColor(30, 41, 59)          # Slate Dark (#1E293B)
    c_muted = RGBColor(100, 116, 139)      # Slate Muted (#64748B)
    
    # --- HEADER / KOP DOKUMEN ---
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    # Col 1: Title
    c1 = header_table.cell(0, 0)
    c1.width = Inches(5.0)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("STANDAR OPERASIONAL PROSEDUR (SOP)")
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = c_primary
    
    p2 = c1.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("SISTEM INFORMASI LEMBUR KARYAWAN (SILEKA)")
    r2.bold = True
    r2.font.size = Pt(15)
    r2.font.color.rgb = c_dark
    
    p3 = c1.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    r3 = p3.add_run("Pabrik Kelapa Sawit (PKS) - PT. Letawa | Format Standar PROS-HCM-02")
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = c_muted
    
    # Col 2: Doc Info Box
    c2 = header_table.cell(0, 1)
    c2.width = Inches(2.2)
    set_cell_background(c2, "F1F5F9")
    set_cell_margins(c2, 80, 80, 100, 100)
    p_info = c2.paragraphs[0]
    p_info.paragraph_format.space_after = Pt(2)
    p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_code = p_info.add_run("Dokumen No: SOP-HCM-2026\nRevisi: 03 (Sistem 3 Aktor)\nStatus: Berlaku Efektif\nKlasifikasi: Internal PKS")
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = c_dark
    
    doc.add_paragraph()
    
    # Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 65)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    
    # --- 1. LATAR BELAKANG & 4 PERMASALAHAN UTAMA SISTEM KERTAS ---
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Latar Belakang & 4 Permasalahan Utama Sistem Kertas")
    r_h1.font.size = Pt(13)
    r_h1.bold = True
    r_h1.font.color.rgb = c_primary
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.add_run(
        "Pada sistem konvensional sebelumnya di Pabrik Kelapa Sawit (PKS), blangko Surat Perintah Lembur (SPL) fisik "
        "sudah memiliki cetakan nama karyawan, nomor employee, dan stasiun kerja. Mandor di lapangan bertugas mengisi secara "
        "manual jam lembur dan jenis pekerjaan satu per satu. Pengelolaan manual ini menimbulkan 4 kendala utama:"
    )
    
    # 4 Masalah Utama Boxes
    issues = [
        ("Pencatatan Berbasis Kertas", "Data lembur (SPL) masih ditulis tangan di atas lembaran kertas, sehingga rawan rusak, basah/terkena minyak di area pabrik, hilang, atau memiliki tulisan tangan yang sulit dibaca saat hendak direkap."),
        ("Beban Rekapitulasi Harian", "Setiap pagi, kertas laporan disetorkan oleh Mandor dan Admin harus mengetik ulang seluruh data fisik tersebut ke Microsoft Excel secara manual, menyita waktu operasional yang sangat berharga."),
        ("Pemborosan Biaya & Ruang Arsip", "Pencetakan blangko SPL setiap hari memakan banyak biaya pembelian kertas & tinta cetak, menghasilkan penumpukan limbah dokumen, serta menghabiskan ruang penyimpanan arsip fisik yang terbatas di kantor."),
        ("Pencarian Arsip Data Sulit", "Jika terjadi komplain selisih perhitungan jam kerja atau audit, Admin kesulitan melacak riwayat lembur karyawan terdahulu karena harus membongkar tumpukan arsip kertas satu per satu.")
    ]
    
    for title, desc in issues:
        t_issue = doc.add_table(rows=1, cols=1)
        t_issue.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_issue.autofit = False
        cell = t_issue.cell(0, 0)
        cell.width = Inches(7.0)
        set_cell_background(cell, "FEF2F2") # Light red
        set_cell_margins(cell, 70, 70, 110, 110)
        
        # Red left border
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '18')
        left.set(qn('w:space'), '0')
        left.set(qn('w:color'), 'DC2626')
        tcBorders.append(left)
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(f"✕ {title}")
        r_t.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = RGBColor(185, 28, 28)
        
        p_d = cell.add_paragraph()
        p_d.paragraph_format.space_after = Pt(0)
        p_d.paragraph_format.line_spacing = 1.15
        r_d = p_d.add_run(desc)
        r_d.font.size = Pt(9)
        r_d.font.color.rgb = RGBColor(51, 65, 85)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(3)
        
    # --- 2. PERBANDINGAN SISTEM KONVENSIONAL VS SILEKA ---
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Perbandingan Sistem Konvensional vs SILEKA")
    r_h2.font.size = Pt(13)
    r_h2.bold = True
    r_h2.font.color.rgb = c_primary
    
    comp_table = doc.add_table(rows=5, cols=3)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_table.autofit = False
    
    headers = ["Aspek Pengelolaan", "Sistem Konvensional (Kertas)", "Sistem Digital (SILEKA)"]
    widths = [Inches(1.8), Inches(2.6), Inches(2.6)]
    
    for i, t_head in enumerate(headers):
        cell = comp_table.cell(0, i)
        cell.width = widths[i]
        set_cell_background(cell, "0D6E55")
        set_cell_margins(cell, 90, 90, 110, 110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(t_head)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    comp_data = [
        ("Pengisian Jam & Tugas", 
         "Mandor mengisi jam lembur dan jenis pekerjaan secara manual dengan pulpen pada 20-35 baris blangko cetak.", 
         "Mandor cukup memakai fitur 'Preset Jam Massal' (1 kali klik terisi serentak) lalu menyesuaikan jenis pekerjaan."),
        ("Rekapitulasi Harian Admin", 
         "Kertas diserahkan tiap pagi; Admin harus mengetik ulang seluruh data fisik ke file Excel satu per satu.", 
         "Data otomatis tersimpan di database saat Mandor klik simpan. Admin tidak perlu lagi mengetik ulang ke Excel."),
        ("Biaya Kertas & Limbah Arsip", 
         "Mencetak lembaran blangko setiap hari, pemborosan kertas, tumpukan berkas kotor, dan lemari arsip penuh.", 
         "100% paperless tersimpan di cloud. Cetak fisik PROS-HCM-02 A4 hanya opsi cadangan jika atasan membutuhkan."),
        ("Pelacakan Riwayat & Komplain", 
         "Admin harus mencari manual membongkar tumpukan map arsip kertas berdebu berbulan-bulan ke belakang.", 
         "Pencarian riwayat lembur instan dalam hitungan detik via filter tanggal dan nama karyawan di sistem.")
    ]
    
    for row_idx, row in enumerate(comp_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row):
            cell = comp_table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = c_dark
            elif col_idx == 1:
                r.font.color.rgb = RGBColor(185, 28, 28)
            else:
                r.font.color.rgb = RGBColor(21, 128, 61)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # --- 3. KETENTUAN SHIFT & ROTASI KERJA ---
    h_shift = doc.add_heading(level=1)
    r_hs = h_shift.add_run("3. Mekanisme Shift & Rotasi Kerja Mingguan")
    r_hs.font.size = Pt(13)
    r_hs.bold = True
    r_hs.font.color.rgb = c_primary
    
    p_shift = doc.add_paragraph()
    p_shift.paragraph_format.space_after = Pt(6)
    p_shift.paragraph_format.line_spacing = 1.15
    p_shift.add_run(
        "Struktur kerja pabrik kelapa sawit diatur berdasarkan regu kerja tetap yang berotasi waktu kerja setiap pekan:\n"
        "• Regu Shift 1 dipimpin oleh Mandor A dengan anggota karyawan yang sudah ditetapkan tetap.\n"
        "• Regu Shift 2 dipimpin oleh Mandor B dengan anggota karyawan yang sudah ditetapkan tetap.\n"
        "• Pergantian Jam Kerja (Rolling): Rotasi dilakukan per minggu. Contohnya: jika minggu ini Shift 1 bertugas pada jam pagi dan Shift 2 bertugas pada jam malam, maka minggu berikutnya Shift 1 berganti bertugas pada jam malam dan Shift 2 pada jam pagi. Istilah 'Shift 1' merujuk pada identitas regu kerja, bukan jam pagi permanen.\n"
        "• Pembuatan SPL Mandor: Di dalam sistem SILEKA, akun Mandor sudah langsung terhubung dengan regu shift-nya masing-masing. Saat Mandor membuka menu 'Buat SPL', sistem secara otomatis memuat anggota regunya tanpa perlu Mandor memilih shift lagi."
    )
    p_shift.runs[0].font.size = Pt(9.5)
    p_shift.runs[0].font.color.rgb = c_dark
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # --- 4. TIGA AKTOR & TUGAS MASING-MASING ---
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("4. Tiga Aktor Utama & Rincian Tugas")
    r_h3.font.size = Pt(13)
    r_h3.bold = True
    r_h3.font.color.rgb = c_primary
    
    actors = [
        ("MANDOR / PENGAWAS LAPANGAN (Shift Supervisor)", "0D6E55", [
            ("Definisi Peran", "Pengawas regu shift operasional di lapangan pabrik (Mandor Shift 1 dan Mandor Shift 2)."),
            ("Tugas 1 (Pembuatan SPL Otomatis)", "Membuka menu 'Buat SPL' dan memilih tanggal. Daftar seluruh karyawan di shift-nya otomatis tampil lengkap dengan nomor employee dan stasiun tetapnya (tanpa perlu memilih shift lagi)."),
            ("Tugas 2 (Pengisian Jam Cepat)", "Menggunakan fitur Preset Bar untuk mengisi jam lembur massal (misal: 17:00 - 21:00) ke seluruh karyawan hanya dengan 1 klik, lalu menyesuaikan jenis pekerjaan."),
            ("Tugas 3 (Status Lembur vs Tidak)", "Menetapkan status Lembur atau Tidak Lembur. Bagi karyawan yang izin, sakit, atau hanya jam kerja normal (HK), status diatur ke 'Tidak Lembur' dan keterangannya diisi di Jenis Pekerjaan."),
            ("Tugas 4 (Simpan Data)", "Menyimpan SPL harian ke dalam server, sehingga seketika dapat dilihat oleh Karyawan di HP masing-masing dan Admin di kantor."),
            ("Tugas 5 (Opsi Cetak PROS-HCM-02)", "Sebagai opsi cadangan: dapat mencetak blangko A4 format resmi PROS-HCM-02 jika atasan menginginkan bukti tanda tangan fisik basah di kertas.")
        ]),
        ("KARYAWAN PELAKSANA / OPERATOR PABRIK (Mill Operator)", "1A56DB", [
            ("Definisi Peran", "Pelaksana operasional teknis di 14 stasiun pabrik (Sterilizer, Thresher, Press, Boiler, dll)."),
            ("Tugas 1 (Akses Mandiri di HP)", "Membuka portal SILEKA melalui browser HP masing-masing tanpa perlu berkumpul atau mengantre di lorong pabrik."),
            ("Tugas 2 (Pemeriksaan Jam Kerja)", "Memeriksa tanggal, jam mulai, jam selesai, dan jenis tugas yang dimasukkan oleh Mandor pada dashboard pribadinya."),
            ("Tugas 3 (Validasi Digital - ACC)", "Mengonfirmasi kebenaran jam lembur dengan menekan tombol 'ACC'. Status lembur langsung berubah menjadi sah ('Disetujui')."),
            ("Tugas 4 (Pengajuan Koreksi - Tolak)", "Jika jam kerja yang diajukan tidak sesuai dengan jam aktual lembur di pabrik, karyawan klik tombol 'Tolak' dan mengetikkan alasan revisi."),
            ("Tugas 5 (Pantau Riwayat)", "Melihat riwayat akumulasi jam lembur bulanan secara transparan kapan saja.")
        ]),
        ("ADMIN SISTEM / PERSONALIA (HCM Staff)", "D97706", [
            ("Definisi Peran", "Pengelola data induk operasional dan pemantau rekapitulasi data lembur pabrik."),
            ("Tugas 1 (Kelola Master Karyawan)", "Mendaftarkan karyawan baru, nomor employee/SAP, penempatan stasiun pabrik permanen, dan penetapan regu shift (Shift 1 atau Shift 2)."),
            ("Tugas 2 (Monitoring Data Tanpa Ketik Ulang)", "Melihat seluruh rekaman data lembur harian yang masuk dari seluruh Mandor secara otomatis di dashboard (menghilangkan tugas mengetik ulang dari kertas ke Excel setiap pagi)."),
            ("Tugas 3 (Penanganan Selisih & Komplain)", "Membuka riwayat data lembur digital secara cepat ketika ada pertanyaan mengenai jam kerja lembur karyawan."),
            ("Tugas 4 (Opsi Cetak & Arsip Cadangan)", "Mencetak formulir cetak PROS-HCM-02 bila diperlukan untuk laporan audit fisik tertentu."),
            ("Tugas 5 (Ekspor Rekapitulasi Payroll)", "Mengunduh rekap total jam lembur yang sudah valid untuk diproses ke dalam perhitungan penggajian bulanan.")
        ])
    ]
    
    for title, color_hex, details in actors:
        t_box = doc.add_table(rows=1, cols=1)
        t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_box.autofit = False
        cell = t_box.cell(0, 0)
        cell.width = Inches(7.0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, 100, 100, 140, 140)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"■ {title}")
        r.bold = True
        r.font.size = Pt(10.5)
        h_r, h_g, h_b = int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
        r.font.color.rgb = RGBColor(h_r, h_g, h_b)
        
        for k, v in details:
            p_sub = cell.add_paragraph()
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.line_spacing = 1.15
            r_k = p_sub.add_run(f"• {k}: ")
            r_k.bold = True
            r_k.font.size = Pt(9)
            r_k.font.color.rgb = c_dark
            
            r_v = p_sub.add_run(v)
            r_v.font.size = Pt(9)
            r_v.font.color.rgb = RGBColor(51, 65, 85)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
    # --- 5. ALUR KERJA OPERASIONAL ---
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("5. Alur Kerja Operasional Harian (Workflow)")
    r_h4.font.size = Pt(13)
    r_h4.bold = True
    r_h4.font.color.rgb = c_primary
    
    steps = [
        ("TAHAP 1: INPUT LEMBUR HARIAN OLEH MANDOR", "Mandor Lapangan", [
            "Mandor login ke SILEKA -> Masuk ke menu 'Buat SPL' dan memilih tanggal operasional.",
            "Sistem otomatis menampilkan seluruh karyawan di shift mandor tersebut (tanpa perlu pilih shift lagi).",
            "Mandor mengisi jam serentak dengan Preset Jam Cepat dan mengisi jenis tugas/pekerjaan.",
            "Untuk karyawan yang tidak lembur (izin/cuti/sakit/HK), toggle diubah ke 'Tidak Lembur' dan alasan ditulis pada kolom jenis pekerjaan.",
            "Mandor klik 'Simpan Semua Data Lembur' -> Data tersimpan di server."
        ]),
        ("TAHAP 2: VALIDASI DIGITAL DI HP OLEH KARYAWAN", "Karyawan Pelaksana", [
            "Karyawan membuka aplikasi SILEKA dari HP masing-masing setelah selesai bekerja.",
            "Di dashboard, karyawan memeriksa rincian lembur (tanggal, jam mulai, jam selesai, jenis tugas).",
            "Jika sesuai fakta kerja di lapangan: Karyawan klik tombol 'ACC' (Status menjadi Disetujui).",
            "Jika jam berbeda: Karyawan klik 'Tolak' dan menulis alasan perbaikan agar diperbaiki oleh Mandor."
        ]),
        ("TAHAP 3: REKAPITULASI OTOMATIS DI SISI ADMIN", "Admin Sistem", [
            "Setiap pagi Admin cukup membuka dashboard SILEKA; seluruh data lembur dari seluruh shift sudah tersaji lengkap dan otomatis terakumulasi.",
            "Admin TIDAK PERLU lagi mengetik ulang dari kertas ke Excel.",
            "Ketika ada selisih perhitungan, Admin dapat langsung mencari riwayat digital secara instan via filter tanggal.",
            "Data siap diekspor untuk keperluan slip gaji lembur bulanan."
        ]),
        ("TAHAP CADANGAN: CETAK FORMULIR FISIK PROS-HCM-02 (OPSIONAL)", "Mandor / Admin", [
            "Jika atasan / pimpinan operasional merasa repot membuka HP untuk ACC digital atau membutuhkan tanda tangan fisik di kertas untuk arsip kantor:",
            "Mandor atau Admin membuka menu 'Cetak SPL'.",
            "Sistem menyajikan format standar PROS-HCM-02 pas 1 halaman A4 portrait bebas dari tampilan web.",
            "Formulir dicetak ke kertas fisik untuk ditandatangani basah oleh Mandor dan Atasan."
        ])
    ]
    
    for title, role, subitems in steps:
        s_table = doc.add_table(rows=1, cols=1)
        s_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        s_table.autofit = False
        cell = s_table.cell(0, 0)
        cell.width = Inches(7.0)
        set_cell_background(cell, "FFFFFF")
        set_cell_margins(cell, 80, 80, 120, 120)
        
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '24')
        left.set(qn('w:space'), '0')
        left.set(qn('w:color'), '0D6E55')
        tcBorders.append(left)
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = c_primary
        
        p_role = cell.add_paragraph()
        p_role.paragraph_format.space_after = Pt(4)
        r_r = p_role.add_run(f"Penanggung Jawab: {role}")
        r_r.italic = True
        r_r.font.size = Pt(8.5)
        r_r.font.color.rgb = c_muted
        
        for idx, item in enumerate(subitems, 1):
            p_i = cell.add_paragraph()
            p_i.paragraph_format.space_after = Pt(2)
            p_i.paragraph_format.line_spacing = 1.15
            r_num = p_i.add_run(f"{idx}. ")
            r_num.bold = True
            r_num.font.size = Pt(9)
            r_num.font.color.rgb = c_dark
            
            r_txt = p_i.add_run(item)
            r_txt.font.size = Pt(9)
            r_txt.font.color.rgb = RGBColor(51, 65, 85)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
    # --- 6. MATRIKS TANGGUNG JAWAB (3 AKTOR) ---
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("6. Matriks Tanggung Jawab (RACI Matrix)")
    r_h5.font.size = Pt(13)
    r_h5.bold = True
    r_h5.font.color.rgb = c_primary
    
    p_raci_desc = doc.add_paragraph()
    p_raci_desc.paragraph_format.space_after = Pt(4)
    p_raci_desc.add_run(
        "Keterangan: R = Responsible (Pelaksana), A = Accountable (Penanggung Jawab Utama), "
        "C = Consulted (Pemberi Masukan), I = Informed (Penerima Laporan/Informasi)."
    )
    p_raci_desc.runs[0].font.size = Pt(8.5)
    p_raci_desc.runs[0].italic = True
    p_raci_desc.runs[0].font.color.rgb = c_muted
    
    raci_table = doc.add_table(rows=6, cols=4)
    raci_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    raci_table.autofit = False
    
    raci_cols = ["Aktivitas / Tahapan Kerja", "Mandor", "Karyawan", "Admin"]
    raci_widths = [Inches(3.4), Inches(1.2), Inches(1.2), Inches(1.2)]
    
    for i, col_name in enumerate(raci_cols):
        c = raci_table.cell(0, i)
        c.width = raci_widths[i]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, 80, 80, 80, 80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(col_name)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    raci_rows = [
        ("Pengelolaan Master Karyawan, Shift & Stasiun", "C", "I", "R / A"),
        ("Penyusunan SPL Harian Regu Shift", "R / A", "I", "I"),
        ("Penetapan Jam Lembur vs Tidak Lembur & Alasan", "R / A", "I", "I"),
        ("Verifikasi & Validasi Digital (ACC / Tolak) di HP", "I", "R / A", "I"),
        ("Rekapitulasi Harian & Penanganan Selisih Riwayat", "C", "I", "R / A")
    ]
    
    for r_idx, row in enumerate(raci_rows, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            c = raci_table.cell(r_idx, c_idx)
            c.width = raci_widths[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 80, 80)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = c_dark
            else:
                r.bold = True
                if "R" in val:
                    r.font.color.rgb = RGBColor(13, 110, 85)
                elif "A" in val:
                    r.font.color.rgb = RGBColor(185, 28, 28)
                else:
                    r.font.color.rgb = RGBColor(100, 116, 139)
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    output_path = os.path.abspath("d:/Projek/web-lembur/public/SOP_DAN_ALUR_SISTEM_SILEKA.docx")
    doc.save(output_path)
    print(f"Updated DOCX created successfully at: {output_path}")

if __name__ == "__main__":
    create_sop_document()
